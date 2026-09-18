"""Excel pack for weekly enablement readout (Power BI-adjacent storytelling)."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font
from openpyxl.utils.dataframe import dataframe_to_rows

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs"


def _sheet(wb: Workbook, title: str, df: pd.DataFrame) -> None:
    ws = wb.create_sheet(title)
    for r in dataframe_to_rows(df, index=False, header=True):
        ws.append(r)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 22


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    funnel = pd.read_csv(OUT / "funnel_activation.csv")
    weekly = pd.read_csv(OUT / "weekly_usage.csv")
    feat = pd.read_csv(OUT / "feature_value_rate.csv")
    ttv = pd.read_csv(OUT / "ttv_by_bu.csv")
    shadow = pd.read_csv(OUT / "shadow_it.csv")
    repeat = pd.read_csv(OUT / "repeat_use.csv")

    _sheet(wb, "Activation funnel", funnel)
    _sheet(wb, "Weekly usage", weekly)
    _sheet(wb, "Feature value rate", feat)
    _sheet(wb, "TTV by BU", ttv)
    _sheet(wb, "Shadow IT proxy", shadow)
    _sheet(wb, "Repeat value use", repeat)

    ws = wb["Weekly usage"]
    chart = LineChart()
    chart.title = "WAU vs value events"
    data = Reference(ws, min_col=3, min_row=1, max_col=4, max_row=ws.max_row)
    cats = Reference(ws, min_col=2, min_row=2, max_row=ws.max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.y_axis.title = "Count"
    chart.height = 8
    chart.width = 18
    ws.add_chart(chart, "G2")

    ws2 = wb["Feature value rate"]
    bar = BarChart()
    bar.title = "Value-event rate by feature"
    bar_data = Reference(ws2, min_col=4, min_row=1, max_row=ws2.max_row)
    bar_cats = Reference(ws2, min_col=1, min_row=2, max_row=ws2.max_row)
    bar.add_data(bar_data, titles_from_data=True)
    bar.set_categories(bar_cats)
    bar.shape = 4
    ws2.add_chart(bar, "F2")

    path = OUT / "genai_adoption_readout.xlsx"
    wb.save(path)
    print(f"Wrote {path}")


if __name__ == "__main__":
    main()
